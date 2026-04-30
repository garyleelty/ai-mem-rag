import uuid
import shutil
from pathlib import Path
from datetime import datetime

import chromadb
from sentence_transformers import SentenceTransformer
from pypdf import PdfReader
from docx import Document

from config import CHROMA_DIR, UPLOAD_DIR, EMBEDDING_MODEL, EMBEDDING_DEVICE, ALLOWED_EXTENSIONS, MAX_UPLOAD_SIZE_BYTES

_chroma_client = None
_embedding_model = None

def get_chroma_client():
    global _chroma_client
    if _chroma_client is None:
        _chroma_client = chromadb.PersistentClient(path=str(CHROMA_DIR))
    return _chroma_client

def get_embedding_model():
    global _embedding_model
    if _embedding_model is None:
        _embedding_model = SentenceTransformer(EMBEDDING_MODEL, device=EMBEDDING_DEVICE)
    return _embedding_model

def validate_file(filename: str, file_size: int):
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"不支持的文件类型: {ext}，仅支持 {', '.join(ALLOWED_EXTENSIONS)}")
    if file_size > MAX_UPLOAD_SIZE_BYTES:
        raise ValueError(f"文件大小超过限制: {MAX_UPLOAD_SIZE_BYTES // 1024 // 1024}MB")

def save_upload(file) -> Path:
    file_id = str(uuid.uuid4())
    ext = Path(file.filename).suffix
    file_path = UPLOAD_DIR / f"{file_id}{ext}"
    with open(file_path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    return file_path

def parse_document(file_path: Path) -> str:
    ext = file_path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext == ".docx":
        return parse_docx(file_path)
    else:
        raise ValueError(f"不支持的文件类型: {ext}")

def parse_pdf(file_path: Path) -> str:
    reader = PdfReader(str(file_path))
    texts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            texts.append(text)
    return "\n\n".join(texts)

def parse_docx(file_path: Path) -> str:
    doc = Document(str(file_path))
    texts = []

    for para in doc.paragraphs:
        if para.text.strip():
            texts.append(para.text)

    for table in doc.tables:
        table_texts = []
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                table_texts.append(" | ".join(row_texts))
        if table_texts:
            texts.append("\n".join(table_texts))

    return "\n\n".join(texts)

def simple_split(text: str, chunk_size: int = 1000) -> list[str]:
    """将文本分割成指定大小的块，优先在段落边界、句子边界或换行符处分割。

    Args:
        text: 要分割的原始文本
        chunk_size: 每个块的目标大小（字符数），默认为1000

    Returns:
        分割后的文本块列表
    """
    if not text:
        return []
    chunks = []
    start = 0
    text_len = len(text)
    while start < text_len:
        end = start + chunk_size
        if end >= text_len:
            chunks.append(text[start:])
            break
        paragraph_end = text.rfind("\n\n", start, end)
        if paragraph_end != -1 and paragraph_end > start + chunk_size * 0.5:
            end = paragraph_end + 2
        else:
            sentence_end = text.rfind(". ", start, end)
            if sentence_end != -1 and sentence_end > start + chunk_size * 0.5:
                end = sentence_end + 2
            else:
                newline_end = text.rfind("\n", start, end)
                if newline_end != -1 and newline_end > start + chunk_size * 0.5:
                    end = newline_end + 1
        chunks.append(text[start:end])
        start = end
    return chunks

def process_file(file, collection_name: str, chunk_size: int = 1000):
    file_path = save_upload(file)
    try:
        text = parse_document(file_path)
        chunks = simple_split(text, chunk_size)
        if not chunks:
            raise ValueError("文档内容为空，无法处理")

        model = get_embedding_model()
        embeddings = model.encode(chunks)

        doc_id = str(uuid.uuid4())[:8]
        client = get_chroma_client()
        collection = client.get_or_create_collection(name=collection_name)

        ids = [f"{doc_id}_{i}" for i in range(len(chunks))]
        metadatas = [
            {
                "source": file.filename,
                "doc_id": doc_id,
                "collection_name": collection_name,
                "chunk_index": i,
                "total_chunks": len(chunks),
                "created_at": datetime.now().isoformat(),
            }
            for i in range(len(chunks))
        ]

        collection.add(
            ids=ids,
            embeddings=embeddings.tolist(),
            documents=chunks,
            metadatas=metadatas,
        )

        return {
            "document_id": doc_id,
            "filename": file.filename,
            "collection_name": collection_name,
            "chunk_count": len(chunks),
            "status": "completed",
        }
    finally:
        if file_path.exists():
            file_path.unlink()

def list_collections():
    client = get_chroma_client()
    collections = client.list_collections()
    result = []
    for col in collections:
        count = col.count()
        result.append({"name": col.name, "count": count})
    return result

def create_collection(name: str):
    client = get_chroma_client()
    client.create_collection(name=name)

def delete_collection(name: str):
    client = get_chroma_client()
    client.delete_collection(name=name)

def list_documents(collection_name: str):
    client = get_chroma_client()
    try:
        collection = client.get_collection(name=collection_name)
    except Exception:
        return {"collection_name": collection_name, "documents": []}

    all_data = collection.get(include=["metadatas"])
    metadatas = all_data.get("metadatas") or []

    docs_map = {}
    for meta in metadatas:
        if not meta:
            continue
        doc_id = meta.get("doc_id")
        if not doc_id:
            continue
        if doc_id not in docs_map:
            docs_map[doc_id] = {
                "id": doc_id,
                "source": meta.get("source", "unknown"),
                "chunk_count": meta.get("total_chunks", 0),
                "created_at": meta.get("created_at"),
            }

    documents = list(docs_map.values())
    documents.sort(key=lambda x: x.get("created_at") or "", reverse=True)
    return {"collection_name": collection_name, "documents": documents}

def get_document_chunks(document_id: str):
    client = get_chroma_client()
    collections = client.list_collections()

    target_collection = None
    for col in collections:
        try:
            result = col.get(
                where={"doc_id": document_id},
                include=["metadatas"],
                limit=1,
            )
            if result and result.get("ids"):
                target_collection = col
                break
        except Exception:
            continue

    if target_collection is None:
        raise ValueError(f"文档不存在: {document_id}")

    result = target_collection.get(
        where={"doc_id": document_id},
        include=["documents", "metadatas"],
    )
    docs = result.get("documents") or []
    metas = result.get("metadatas") or []
    ids = result.get("ids") or []

    chunks = []
    for i, chunk_id in enumerate(ids):
        meta = metas[i] if i < len(metas) else {}
        chunks.append({
            "chunk_id": chunk_id,
            "index": meta.get("chunk_index", i),
            "content": docs[i] if i < len(docs) else "",
        })
    chunks.sort(key=lambda x: x["index"])

    source = metas[0].get("source", "unknown") if metas else "unknown"
    return {
        "document_id": document_id,
        "document_name": source,
        "chunks": chunks,
    }

def delete_document(document_id: str):
    client = get_chroma_client()
    collections = client.list_collections()

    target_collection = None
    for col in collections:
        try:
            result = col.get(where={"doc_id": document_id}, limit=1)
            if result and result.get("ids"):
                target_collection = col
                break
        except Exception:
            continue

    if target_collection is None:
        raise ValueError(f"文档不存在: {document_id}")

    result = target_collection.get(where={"doc_id": document_id})
    target_collection.delete(ids=result["ids"])
    return True
