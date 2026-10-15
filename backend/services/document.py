import uuid
import shutil
from pathlib import Path
from datetime import datetime

from pypdf import PdfReader
from docx import Document

from config import UPLOAD_DIR, ALLOWED_EXTENSIONS, MAX_UPLOAD_SIZE_BYTES
from services.chroma import get_chroma_client
from services.embedding import get_embedding_model


def validate_file(filename: str, file_size: int):
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"不支持的文件类型: {ext}，仅支持 {', '.join(ALLOWED_EXTENSIONS)}")
    if file_size > MAX_UPLOAD_SIZE_BYTES:
        raise ValueError(f"文件大小超过限制: {MAX_UPLOAD_SIZE_BYTES // 1024 // 1024}MB")


def save_upload(file) -> Path:
    file_id = str(uuid.uuid4())
    ext = Path(file.filename).suffix
    file_path = UPLOAD_DIR / f"{file_id}{ext}"
    with open(file_path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    return file_path


def parse_document(file_path: Path) -> str:
    ext = file_path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext == ".docx":
        return parse_docx(file_path)
    else:
        raise ValueError(f"不支持的文件类型: {ext}")


def parse_pdf(file_path: Path) -> str:
    reader = PdfReader(str(file_path))
    texts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            texts.append(text)
    return "\n\n".join(texts)


def parse_docx(file_path: Path) -> str:
    doc = Document(str(file_path))
    texts = []

    for para in doc.paragraphs:
        if para.text.strip():
            texts.append(para.text)

    for table in doc.tables:
        table_texts = []
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                table_texts.append(" | ".join(row_texts))
        if table_texts:
            texts.append("\n".join(table_texts))

    return "\n\n".join(texts)


def simple_split(text: str, chunk_size: int = 1000) -> list[str]:
    if not text:
        return []
    chunks = []
    start = 0
    text_len = len(text)
    while start < text_len:
        end = start + chunk_size
        if end >= text_len:
            chunks.append(text[start:])
            break
        paragraph_end = text.rfind("\n\n", start, end)
        if paragraph_end != -1 and paragraph_end > start + chunk_size * 0.5:
            end = paragraph_end + 2
        else:
            sentence_end = text.rfind(". ", start, end)
            if sentence_end != -1 and sentence_end > start + chunk_size * 0.5:
                end = sentence_end + 2
            else:
                newline_end = text.rfind("\n", start, end)
                if newline_end != -1 and newline_end > start + chunk_size * 0.5:
                    end = newline_end + 1
        chunks.append(text[start:end])
        start = end
    return chunks


def process_file(file, collection_name: str, chunk_size: int = 1000):
    file_path = save_upload(file)
    try:
        text = parse_document(file_path)
        chunks = simple_split(text, chunk_size)
        if not chunks:
            raise ValueError("文档内容为空，无法处理")

        model = get_embedding_model()
        embeddings = model.encode(chunks)

        doc_id = str(uuid.uuid4())[:8]
        client = get_chroma_client()
        collection = client.get_or_create_collection(name=collection_name)

        ids = [f"{doc_id}_{i}" for i in range(len(chunks))]
        metadatas = [
            {
                "source": file.filename,
                "doc_id": doc_id,
                "collection_name": collection_name,
                "chunk_index": i,
                "total_chunks": len(chunks),
                "created_at": datetime.now().isoformat(),
            }
            for i in range(len(chunks))
        ]

        collection.add(
            ids=ids,
            embeddings=embeddings.tolist(),
            documents=chunks,
            metadatas=metadatas,
        )

        return {
            "document_id": doc_id,
            "filename": file.filename,
            "collection_name": collection_name,
            "chunk_count": len(chunks),
            "status": "completed",
        }
    finally:
        if file_path.exists():
            file_path.unlink()
